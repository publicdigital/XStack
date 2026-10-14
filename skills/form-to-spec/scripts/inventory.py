#!/usr/bin/env python3
"""Inventory a government form before turning it into a form spec.

Usage:
  inventory.py FORM [--render DIR] [--dpi 110] [--text-out FILE]

FORM can be a PDF (fillable or flat), an image (scan or photo) or a Word
document (.doc, .docx, .rtf, .odt). Prints JSON to stdout:

  file, kind, pages, metadata (title, producer, creator, language),
  acroform fields (name, type, tooltip, options, page), text per page,
  whether the PDF looks scanned, which pages to read visually, and
  hints (possible legal citations, fees, signature and oath wording).

Hints are pointers for the agent to check against the form itself. They are
never facts to copy into the spec.

Uses pypdf (text and fields) and pypdfium2 (rendering) when installed, and
textutil (macOS) or python-docx for Word files. Without them, it says what it
couldn't do, and the agent reads the file visually instead.
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".heic", ".tif", ".tiff", ".webp", ".gif", ".bmp"}
WORD_EXTS = {".doc", ".docx", ".rtf", ".odt"}

# Fewer characters than this on a PDF page means it is probably an image of text.
SCANNED_CHARS_PER_PAGE = 40

HINT_PATTERNS = {
    "legal_citation": [
        r"\b[A-Z][A-Za-z ,'()-]{2,80}\b(?:Act|Ordinance|Decree|Code|Law)\b(?:,? *(?:19|20)\d\d)?(?:[-–/ ]\d+)?",
        r"\b(?:Regulations|Rules|Order)\b,? *(?:19|20)\d\d",
        r"(?i:\b(?:First|Second|Third|Fourth|Fifth|Sixth|\d+(?:st|nd|rd|th)?)\s+Schedule\b)",
        r"\bSchedule\s+[0-9IVX]+\b",
        r"(?i:\b(?:section|sec\.|s\.|regulation|reg\.|rule|article|art\.|paragraph|para\.)\s*\d+[A-Za-z]?(?:\(\w+\))*)",
        r"\bAct\s+(?:No\.?\s*)?\d{1,4}[-–/ ]\d+",
        r"\b(?:S\.I\.|SI|Statutory Instrument)\s*(?:No\.?\s*)?\d{4}/\d+",
        r"\bCap(?:ter)?\.?\s*\d+[A-Z]?\b",
    ],
    "fee": [
        r"(?:[$£€₹₦]|BDS\$|BBD|GBP|USD|KES|KSh)\s?\d[\d,]*(?:\.\d\d)?",
        r"\b(?:fee|fees|payable|payment|receipt|stamp duty|revenue stamp)\b",
    ],
    "signature_or_oath": [
        r"\bsign(?:ed|ature)?\b",
        r"\bwitness(?:ed)?\b",
        r"\b(?:oath|affirm(?:ation)?|sworn|swear|solemnly)\b",
        r"\b(?:declar(?:e|ation)|statutory declaration)\b",
        r"\b(?:notary|notarised|notarized|commissioner for oaths|justice of the peace|J\.P\.)\b",
        r"\b(?:countersign\w*|attest\w*|seal)\b",
    ],
    "official_use": [
        r"\bfor (?:official|office) use(?: only)?\b",
        r"\bofficial use\b",
        r"\bdo not write (?:in|below)\b",
    ],
    "form_number": [
        r"\b(?:Form|FORM)\s+(?:No\.?\s*)?[A-Z0-9][A-Z0-9./-]*\b",
    ],
}


def _hints(text):
    # Blank lines on paper forms are underscores, which are word characters and break \b.
    text = text.replace("_", " ")
    found = {}
    for key, patterns in HINT_PATTERNS.items():
        hits = []
        for p in patterns:
            for m in re.finditer(p, text, flags=0 if key in ("legal_citation", "form_number") else re.IGNORECASE):
                s = " ".join(m.group(0).split())
                if s.lower() not in (h.lower() for h in hits):
                    hits.append(s)
        found[key] = hits[:40]
    return found


def _missing(what, why):
    return {"skipped": what, "because": why}


def inventory_pdf(path, render_dir, dpi):
    out = {"kind": "pdf", "acroform": None, "pages": None, "metadata": {}, "text": [], "notes": []}
    try:
        import pypdf  # noqa: F401
    except ImportError:
        out["notes"].append(_missing("text and form fields", "pypdf isn't installed (pip install pypdf)"))
        out["read_visually"] = "all"
        return out

    from pypdf import PdfReader

    reader = PdfReader(path)
    out["pages"] = len(reader.pages)
    meta = reader.metadata or {}
    out["metadata"] = {
        "title": meta.get("/Title"),
        "author": meta.get("/Author"),
        "producer": meta.get("/Producer"),
        "creator": meta.get("/Creator"),
        "language": str(reader.trailer["/Root"].get("/Lang")) if reader.trailer["/Root"].get("/Lang") else None,
    }

    # Fillable fields, with the page each widget sits on.
    page_of = {}
    for i, page in enumerate(reader.pages, start=1):
        for annot in page.get("/Annots") or []:
            a = annot.get_object()
            if a.get("/Subtype") == "/Widget":
                name = a.get("/T") or (a.get("/Parent").get_object().get("/T") if a.get("/Parent") else None)
                if name is not None:
                    page_of.setdefault(str(name), i)
    fields = reader.get_fields() or {}
    if fields:
        rows = []
        for name, f in fields.items():
            ftype = f.get("/FT")
            kind = {"/Tx": "text", "/Btn": "button", "/Ch": "choice", "/Sig": "signature"}.get(ftype, str(ftype))
            if ftype == "/Btn":
                states = f.get("/_States_", [])
                kind = "radio" if len([s for s in states if s != "/Off"]) > 1 else "checkbox"
            short = name.split(".")[-1]
            rows.append({
                "name": name,
                "type": kind,
                "tooltip": f.get("/TU"),
                "options": [str(o) for o in (f.get("/Opt") or f.get("/_States_") or [])] or None,
                "max_length": f.get("/MaxLen"),
                "page": page_of.get(name) or page_of.get(short),
            })
        by_type = {}
        for r in rows:
            by_type[r["type"]] = by_type.get(r["type"], 0) + 1
        out["acroform"] = {"count": len(rows), "by_type": by_type, "fields": rows}

    for i, page in enumerate(reader.pages, start=1):
        try:
            txt = page.extract_text(extraction_mode="layout")
        except Exception:  # older pypdf, or odd encodings
            txt = page.extract_text() or ""
        # Collapse runs of blank lines, keep layout spacing within lines.
        txt = re.sub(r"\n\s*\n\s*\n+", "\n\n", txt or "").rstrip()
        out["text"].append({"page": i, "chars": len(txt.strip()), "text": txt})

    low = [p["page"] for p in out["text"] if p["chars"] < SCANNED_CHARS_PER_PAGE]
    out["looks_scanned"] = len(low) == out["pages"]
    out["read_visually"] = low if low else []

    if render_dir:
        out["rendered"] = render_pdf(path, render_dir, dpi, out["notes"])
    return out


def render_pdf(path, render_dir, dpi, notes):
    try:
        import pypdfium2 as pdfium
    except ImportError:
        notes.append(_missing("rendering pages to images", "pypdfium2 isn't installed (pip install pypdfium2) – read the PDF pages directly instead"))
        return []
    os.makedirs(render_dir, exist_ok=True)
    pdf = pdfium.PdfDocument(path)
    stem = os.path.splitext(os.path.basename(path))[0]
    files = []
    for i in range(len(pdf)):
        img = pdf[i].render(scale=dpi / 72).to_pil()
        f = os.path.join(render_dir, f"{stem}-p{i + 1}.png")
        img.save(f)
        files.append(f)
    return files


def inventory_word(path):
    out = {"kind": "word", "pages": None, "metadata": {}, "text": [], "notes": [], "acroform": None}
    text = None
    if shutil.which("textutil"):
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", path], capture_output=True, text=True)
        if r.returncode == 0:
            text = r.stdout
    if text is None and path.lower().endswith(".docx"):
        try:
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        except ImportError:
            pass
    if text is None and shutil.which("soffice"):
        outdir = os.path.join(os.path.dirname(os.path.abspath(path)), ".inventory-tmp")
        subprocess.run(["soffice", "--headless", "--convert-to", "txt:Text", "--outdir", outdir, path], capture_output=True)
        cand = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".txt")
        if os.path.exists(cand):
            text = open(cand, encoding="utf-8", errors="replace").read()
    if text is None:
        out["notes"].append(_missing("text", "no converter found (textutil, python-docx or LibreOffice) – ask for a PDF export"))
        out["read_visually"] = "all"
        return out
    out["notes"].append("Word export loses tick boxes, signature lines and table layout. Ask for the PDF or a scan if layout matters.")
    out["text"].append({"page": None, "chars": len(text.strip()), "text": text.rstrip()})
    out["read_visually"] = []
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("form")
    ap.add_argument("--render", metavar="DIR", help="render PDF pages to PNGs in DIR for visual reading")
    ap.add_argument("--dpi", type=int, default=110)
    ap.add_argument("--text-out", metavar="FILE", help="also write the extracted text to FILE")
    args = ap.parse_args()

    path = args.form
    if not os.path.exists(path):
        sys.exit(f"No such file: {path}")
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        out = inventory_pdf(path, args.render, args.dpi)
    elif ext in IMAGE_EXTS:
        out = {"kind": "image", "pages": 1, "metadata": {}, "text": [], "acroform": None,
               "looks_scanned": True, "read_visually": "all",
               "notes": ["Image input: read it visually. Note anything you can't read as illegible rather than guessing."]}
    elif ext in WORD_EXTS:
        out = inventory_word(path)
    else:
        sys.exit(f"Unsupported file type: {ext}. Use a PDF, an image or a Word document.")

    all_text = "\n".join(p["text"] for p in out["text"])
    out["hints"] = _hints(all_text) if all_text.strip() else None
    out = {"file": os.path.abspath(path), **out}

    if args.text_out:
        with open(args.text_out, "w", encoding="utf-8") as fh:
            for p in out["text"]:
                fh.write(f"\n===== page {p['page']} =====\n{p['text']}\n")

    json.dump(out, sys.stdout, indent=2, default=str, ensure_ascii=False)
    print()


if __name__ == "__main__":
    main()
